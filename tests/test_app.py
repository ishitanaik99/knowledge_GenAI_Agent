"""Run:  pytest -q      (no API key needed: the LLM is simulated)"""
import io
import json

import httpx
import pytest
from fastapi.testclient import TestClient

from app import config, llm, store
from app.llm import AnthropicLLM, LLMReply, OpenAICompatLLM


class ScriptedLLM:
    """Plays back a fixed list of tool calls, like a real model would."""
    provider = "test"

    def __init__(self, script, enrich=None):
        self.script, self.calls, self.enrich = list(script), [], enrich or {}

    def chat(self, system, transcript, tools, force_tools=True, max_tokens=1024):
        self.calls.append([dict(t) for t in transcript])
        calls = self.script.pop(0)
        return LLMReply(text="", tool_calls=[{"id": f"c{i}", "name": n, "input": a} for i, (n, a) in enumerate(calls)],
                        input_tokens=100, output_tokens=20)

    def complete_json(self, prompt, max_tokens=600):
        return self.enrich


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    monkeypatch.setattr(config, "ACCESS_CODE", "")
    store.reset_singletons()
    llm.set_llm_for_tests(None)
    monkeypatch.setattr(config, "LLM_API_KEY", "")
    with TestClient(app_module().app) as c:
        yield c
    llm.set_llm_for_tests(None)


def app_module():
    from app import main
    return main


def sse(resp):
    return [json.loads(line[6:]) for line in resp.text.splitlines() if line.startswith("data: ")]


def test_health_and_samples(client):
    assert client.get("/api/health").json() == {"ok": True}
    assert client.get("/api/config").json()["llm_ready"] is False
    docs = client.get("/api/documents").json()["documents"]
    assert len(docs) == 8 and sum(d["section_count"] for d in docs) == 27
    assert client.get("/").status_code == 200


def test_agent_answers_with_verified_citation(client):
    fake = ScriptedLLM([
        [("search_sops", {"query": "CAB meeting"})],
        [("read_section", {"section_id": "SOP-DLV-021-S02"})],
        [("submit_answer", {"status": "answered", "citations": ["SOP-DLV-021-S02", "SOP-CARD-007-S01"],
                            "answer": "The CAB meets every Tuesday and Thursday [SOP-DLV-021-S02]. Fake [SOP-CARD-007-S01]."})],
    ])
    llm.set_llm_for_tests(fake)
    events = sse(client.post("/api/chat", json={"question": "When does the CAB meet?"}))
    steps = [e for e in events if e["type"] == "step"]
    final = events[-1]["result"]
    assert [s["tool"] for s in steps] == ["search_sops", "read_section"]
    assert final["status"] == "answered"
    assert [c["section_id"] for c in final["citations"]] == ["SOP-DLV-021-S02"]       # unread citation removed
    assert "[SOP-CARD-007-S01]" not in final["answer"] and "[SOP-DLV-021-S02]" in final["answer"]
    # tool results were sent back to the model in the next call
    assert fake.calls[1][-1]["role"] == "tool_results"


def test_guardrail_escalates_unverified_answer(client):
    llm.set_llm_for_tests(ScriptedLLM([
        [("search_sops", {"query": "canteen budget"})],
        [("submit_answer", {"status": "answered", "answer": "The CFO approves it.", "citations": []})],
    ]))
    final = sse(client.post("/api/chat", json={"question": "Who approves the canteen budget?"}))[-1]["result"]
    assert final["status"] == "escalated" and final["ticket"].startswith("SME-")
    assert "CFO" not in final["answer"]
    assert client.get("/api/tickets").json()["tickets"][0]["status"] == "open"


def test_sme_answer_becomes_searchable(client):
    llm.set_llm_for_tests(ScriptedLLM([
        [("search_sops", {"query": "canteen budget"})],
        [("escalate_to_sme", {"question": "Who approves the canteen budget?", "reason": "No match"})],
        [("submit_answer", {"status": "escalated", "answer": "Sent to an expert.", "citations": []})],
    ]))
    final = sse(client.post("/api/chat", json={"question": "Who approves the canteen budget?"}))[-1]["result"]
    ticket = final["ticket"]
    r = client.post(f"/api/tickets/{ticket}/answer", json={"answer": "The facilities head approves the canteen budget."})
    assert r.status_code == 200
    assert client.post(f"/api/tickets/{ticket}/answer", json={"answer": "again"}).status_code == 409
    results, label = store.library().search("who approves canteen budget")
    assert results[0]["doc_id"] == "SME-ANSWERS" and label in ("Strong", "Good")


def test_retrieval_only_mode(client):
    ok = sse(client.post("/api/chat", json={"question": "When does the CAB meet?"}))[-1]["result"]
    assert ok["status"] == "answered" and ok["mode"] == "retrieval-only" and ok["citations"]
    no = sse(client.post("/api/chat", json={"question": "Who won the cricket match yesterday?"}))[-1]["result"]
    assert no["status"] == "escalated"


def test_upload_parse_and_delete(client):
    from docx import Document
    d = Document()
    d.add_heading("Vendor Onboarding", 0)
    d.add_heading("Documents needed", 1)
    d.add_paragraph("Vendor submits PAN, GST certificate and a cancelled cheque.")
    buf = io.BytesIO()
    d.save(buf)
    files = [("files", ("vendor_onboarding.docx", buf.getvalue())),
             ("files", ("travel.txt", b"Flights under 4 hours are economy class. Submit claims within 15 days.")),
             ("files", ("virus.exe", b"MZ"))]
    r = client.post("/api/documents", files=files).json()
    assert {d["doc_id"] for d in r["added"]} == {"VENDOR-ONBOARDING", "TRAVEL"}
    assert len(r["errors"]) == 1
    results, _ = store.library().search("GST certificate vendor")
    assert results[0]["doc_id"] == "VENDOR-ONBOARDING"
    assert client.delete("/api/documents/TRAVEL").json() == {"ok": True}
    assert client.delete("/api/documents/TRAVEL").status_code == 404
    client.post("/api/documents/reset")
    assert len(client.get("/api/documents").json()["documents"]) == 8


def test_upload_enrichment(client):
    llm.set_llm_for_tests(ScriptedLLM([], enrich={"title": "Vendor Onboarding Policy", "summary": "How vendors join.",
                                                  "questions": ["What documents does a vendor need?"]}))
    client.post("/api/documents", files=[("files", ("vendor_policy.txt", b"Vendors submit PAN and GST certificate."))])
    doc = client.get("/api/documents/VENDOR-POLICY").json()
    assert doc["title"] == "Vendor Onboarding Policy" and doc["questions"]
    sugg = client.get("/api/suggestions").json()["suggestions"]
    assert sugg[0]["question"] == "What documents does a vendor need?"


def test_access_code(client, monkeypatch):
    monkeypatch.setattr(config, "ACCESS_CODE", "pilot2026")
    assert client.get("/api/documents").status_code == 401
    assert client.get("/api/documents", headers={"X-Access-Code": "pilot2026"}).status_code == 200
    assert client.get("/api/config").json()["access_required"] is True


def test_evaluation_and_stats(client):
    ev = client.get("/api/evaluate").json()
    assert ev["summary"]["overall"] == 100 and ev["summary"]["questions"] == 20
    client.post("/api/feedback", json={"question": "q", "rating": "helpful"})
    s = client.get("/api/stats").json()
    assert s["helpful_pct"] == 100 and s["documents"] == 8


# ---------- provider adapters, checked against the real request/response formats ----------
TRANSCRIPT = [
    {"role": "user", "text": "When does the CAB meet?"},
    {"role": "assistant", "text": "", "tool_calls": [{"id": "t1", "name": "search_sops", "input": {"query": "CAB"}}]},
    {"role": "tool_results", "results": [{"id": "t1", "name": "search_sops", "content": "{}"}]},
]
TOOLS = [{"name": "search_sops", "description": "d", "parameters": {"type": "object", "properties": {}}}]


def test_anthropic_adapter():
    seen = {}

    def handler(request):
        seen["url"], seen["headers"], seen["body"] = str(request.url), request.headers, json.loads(request.content)
        return httpx.Response(200, json={"content": [{"type": "text", "text": "ok"}, {"type": "tool_use", "id": "t2",
                              "name": "read_section", "input": {"section_id": "X-S01"}}],
                              "usage": {"input_tokens": 10, "output_tokens": 5}, "stop_reason": "tool_use"})

    a = AnthropicLLM("sk-ant-x", "claude-haiku-5-5", "https://api.anthropic.com", transport=httpx.MockTransport(handler))
    reply = a.chat("sys", TRANSCRIPT, TOOLS)
    assert seen["url"].endswith("/v1/messages") and seen["headers"]["x-api-key"] == "sk-ant-x"
    b = seen["body"]
    assert b["tool_choice"] == {"type": "any"} and b["tools"][0]["input_schema"] == {"type": "object", "properties": {}}
    assert b["messages"][1]["content"][0]["type"] == "tool_use"
    assert b["messages"][2]["content"][0] == {"type": "tool_result", "tool_use_id": "t1", "content": "{}", "is_error": False}
    assert reply.tool_calls == [{"id": "t2", "name": "read_section", "input": {"section_id": "X-S01"}}]
    assert reply.text == "ok" and reply.input_tokens == 10


def test_openai_adapter():
    seen = {}

    def handler(request):
        seen["url"], seen["body"] = str(request.url), json.loads(request.content)
        return httpx.Response(200, json={"choices": [{"message": {"content": None, "tool_calls": [
            {"id": "t2", "type": "function", "function": {"name": "read_section", "arguments": "{\"section_id\": \"X-S01\"}"}}]}}],
            "usage": {"prompt_tokens": 12, "completion_tokens": 4}})

    o = OpenAICompatLLM("sk-x", "gpt-4o-mini", "https://api.openai.com/v1", transport=httpx.MockTransport(handler))
    reply = o.chat("sys", TRANSCRIPT, TOOLS)
    b = seen["body"]
    assert seen["url"].endswith("/chat/completions") and b["tool_choice"] == "required"
    assert b["messages"][0]["role"] == "system"
    assert b["messages"][2]["tool_calls"][0]["function"] == {"name": "search_sops", "arguments": "{\"query\": \"CAB\"}"}
    assert b["messages"][3] == {"role": "tool", "tool_call_id": "t1", "content": "{}"}
    assert reply.tool_calls[0]["input"] == {"section_id": "X-S01"} and reply.output_tokens == 4


def test_adapter_error_message():
    a = AnthropicLLM("bad", "m", "https://api.anthropic.com",
                     transport=httpx.MockTransport(lambda r: httpx.Response(401, json={"error": {"message": "invalid x-api-key"}})))
    with pytest.raises(llm.LLMError, match="API key was rejected"):
        a.chat("s", TRANSCRIPT, TOOLS)
