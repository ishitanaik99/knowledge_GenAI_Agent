"""Turn an uploaded file into a document record plus sections (the pieces the agent searches and cites)."""
import io
import re
from datetime import datetime, timezone
from pathlib import Path

from . import config


class IngestError(ValueError):
    pass


def make_doc_id(name: str) -> str:
    stem = Path(name).stem
    slug = re.sub(r"[^A-Za-z0-9]+", "-", stem).strip("-").upper()
    return (slug or "DOC")[:40]


def _clean(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def _pretty_title(name: str) -> str:
    stem = re.sub(r"[_\-]+", " ", Path(name).stem).strip()
    return stem[:1].upper() + stem[1:] if stem else "Untitled document"


def _finish(doc: dict, parts: list, kind: str, filename: str):
    parts = [(h.strip()[:120] or "Section", _clean(t)) for h, t in parts if _clean(t)]
    if not parts:
        raise IngestError(f"No readable text found in {filename}.")
    doc.update(source=filename, kind=kind, uploaded_at=datetime.now(timezone.utc).isoformat(timespec="seconds"),
               summary=doc.get("summary", ""), questions=doc.get("questions", []))
    doc.setdefault("owner", "")
    doc.setdefault("version", "")
    doc.setdefault("updated", "")
    sections = [{"id": f"{doc['doc_id']}-S{i:02d}", "doc_id": doc["doc_id"], "title": doc["title"],
                 "section": h, "text": t} for i, (h, t) in enumerate(parts, start=1)]
    doc["section_count"] = len(sections)
    doc["word_count"] = sum(len(s["text"].split()) for s in sections)
    return doc, sections


def parse_markdown(text: str, filename: str, kind: str):
    """'# DOC-ID | Title' (optional), 'Owner: .. | Version: .. | Updated: ..' (optional), '## Section' blocks."""
    doc = {"doc_id": make_doc_id(filename), "title": _pretty_title(filename), "title_auto": True}
    parts, heading, buf = [], None, []
    for line in text.splitlines():
        if re.match(r"^#\s", line):
            head = line[2:].strip()
            if "|" in head:
                did, title = [p.strip() for p in head.split("|", 1)]
                doc["doc_id"], doc["title"], doc["title_auto"] = make_doc_id(did), title, False
            elif head:
                doc["title"], doc["title_auto"] = head, False
        elif line.lower().startswith("owner:"):
            for part in line.split("|"):
                if ":" in part:
                    k, v = part.split(":", 1)
                    doc[k.strip().lower()] = v.strip()
        elif re.match(r"^#{2,3}\s", line):
            if heading is not None or "".join(buf).strip():
                parts.append((heading or "Introduction", " ".join(buf)))
            heading, buf = line.lstrip("#").strip(), []
        else:
            buf.append(line)
    parts.append((heading or "Full document", " ".join(buf)))
    if len(parts) == 1 and len(parts[0][1].split()) > config.PDF_CHUNK_WORDS * 2:   # long plain text: chunk it
        words = parts[0][1].split()
        size = config.PDF_CHUNK_WORDS
        parts = [(f"Part {i // size + 1}", " ".join(words[i:i + size])) for i in range(0, len(words), size)]
    return _finish(doc, parts, kind, filename)


def parse_docx(data: bytes, filename: str, kind: str):
    from docx import Document
    try:
        d = Document(io.BytesIO(data))
    except Exception as e:
        raise IngestError(f"Couldn't open {filename} as a Word file.") from e
    doc = {"doc_id": make_doc_id(filename), "title": _pretty_title(filename), "title_auto": True}
    parts, heading, buf = [], "Introduction", []
    for p in d.paragraphs:
        style = (p.style.name or "").lower() if p.style is not None else ""
        txt = p.text.strip()
        if not txt:
            continue
        if style == "title":
            doc["title"], doc["title_auto"] = txt, False
        elif style.startswith("heading"):
            if buf:
                parts.append((heading, " ".join(buf)))
            heading, buf = txt, []
        else:
            buf.append(txt)
    for table in d.tables:                       # keep table text too (fees, timelines ...)
        rows = [" | ".join(c.text.strip() for c in r.cells) for r in table.rows]
        buf.append(" ; ".join(rows))
    if buf:
        parts.append((heading, " ".join(buf)))
    return _finish(doc, parts, kind, filename)


def parse_pdf(data: bytes, filename: str, kind: str):
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [(n, page.extract_text() or "") for n, page in enumerate(reader.pages, start=1)]
    except Exception as e:
        raise IngestError(f"Couldn't read {filename} as a PDF.") from e
    doc = {"doc_id": make_doc_id(filename), "title": _pretty_title(filename), "title_auto": True}
    parts, words, start = [], [], 1
    for n, text in pages:
        for w in text.split():
            if not words:
                start = n
            words.append(w)
            if len(words) >= config.PDF_CHUNK_WORDS:
                parts.append((f"Page {start}", " ".join(words)))
                words = []
    if words:
        parts.append((f"Page {start}", " ".join(words)))
    if not parts:
        raise IngestError(f"{filename} has no selectable text (it may be a scanned image).")
    return _finish(doc, parts, kind, filename)


def parse_file(filename: str, data: bytes, kind: str = "upload"):
    ext = Path(filename).suffix.lower()
    if ext not in config.ALLOWED_TYPES:
        raise IngestError(f"{filename}: only PDF, Word (.docx), Markdown and text files are supported.")
    if ext == ".pdf":
        return parse_pdf(data, filename, kind)
    if ext == ".docx":
        return parse_docx(data, filename, kind)
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    return parse_markdown(text, filename, kind)
